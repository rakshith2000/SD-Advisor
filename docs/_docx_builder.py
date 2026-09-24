"""Shared Word-document builder used by the docs generators.

Wraps python-docx with the handful of primitives these documents need -
headings, shaded code blocks, callouts, tables, checklists - so the generator
scripts contain content and nothing else.
"""

import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Palette (matches the application's documented colour roles)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)
CRITICAL = RGBColor(0xB0, 0x2E, 0x2E)
GOOD = RGBColor(0x00, 0x63, 0x00)

SHADE_CODE = 'F4F4F1'
SHADE_INFO = 'EAF2FC'
SHADE_WARN = 'FDF4E3'
SHADE_CRIT = 'FBECEC'
SHADE_HEAD = 'E8E8E3'


# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------

def shade(element, hex_fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    element.append(shd)


def para_borders(paragraph, hex_color='D8D8D2', left_only=False, size=6):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    sides = ['left'] if left_only else ['top', 'left', 'bottom', 'right']
    for side in sides:
        edge = OxmlElement(f'w:{side}')
        edge.set(qn('w:val'), 'single')
        edge.set(qn('w:sz'), str(size * 4 if left_only else size))
        edge.set(qn('w:space'), '6')
        edge.set(qn('w:color'), hex_color)
        borders.append(edge)
    pPr.append(borders)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement('w:keepNext')
    pPr.append(el)


class Guide:
    def __init__(self, footer_label='Aged Ticket Advisor',
                 header_label='Internal - Service Desk Engineering'):
        self.footer_label = footer_label
        self.header_label = header_label
        self.doc = Document()
        self._setup_styles()
        self._setup_page()

    # -- document chrome ---------------------------------------------------

    def _setup_styles(self):
        styles = self.doc.styles

        normal = styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.12

        for name, size, color, before, after in [
            ('Heading 1', 19, ACCENT, 20, 8),
            ('Heading 2', 14.5, INK, 16, 6),
            ('Heading 3', 12, INK, 12, 4),
            ('Heading 4', 10.5, INK2, 10, 3),
        ]:
            st = styles[name]
            st.font.name = 'Calibri'
            st.font.size = Pt(size)
            st.font.color.rgb = color
            st.font.bold = True
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True

        styles['Title'].font.name = 'Calibri'
        styles['Title'].font.size = Pt(30)
        styles['Title'].font.color.rgb = INK

    def _setup_page(self):
        section = self.doc.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f'{self.footer_label}     |     Page ')
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        p._p.append(fld)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(self.header_label)
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = MUTED

    # -- content primitives ------------------------------------------------

    def h1(self, text, page_break=True):
        if page_break:
            self.doc.add_page_break()
        self.doc.add_heading(text, 1)

    def h2(self, text):
        self.doc.add_heading(text, 2)

    def h3(self, text):
        self.doc.add_heading(text, 3)

    def h4(self, text):
        self.doc.add_heading(text, 4)

    def p(self, text='', bold=False, italic=False, size=10.5, color=None, after=None):
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def rich(self, parts, after=None):
        """parts: list of (text, {'b':True,'i':True,'code':True,'color':RGB})."""
        para = self.doc.add_paragraph()
        for text, fmt in parts:
            run = para.add_run(text)
            run.bold = fmt.get('b', False)
            run.italic = fmt.get('i', False)
            if fmt.get('code'):
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                shade(run._r.get_or_add_rPr(), SHADE_CODE)
            else:
                run.font.size = Pt(fmt.get('size', 10.5))
            if fmt.get('color'):
                run.font.color.rgb = fmt['color']
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def bullet(self, text, level=0, bold_prefix=None):
        para = self.doc.add_paragraph(style='List Bullet')
        para.paragraph_format.left_indent = Cm(0.7 + 0.6 * level)
        para.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            run = para.add_run(bold_prefix)
            run.bold = True
            run.font.size = Pt(10.5)
        run = para.add_run(text)
        run.font.size = Pt(10.5)
        return para

    def numbered(self, text, bold_prefix=None):
        para = self.doc.add_paragraph(style='List Number')
        para.paragraph_format.left_indent = Cm(0.7)
        para.paragraph_format.space_after = Pt(3)
        if bold_prefix:
            run = para.add_run(bold_prefix)
            run.bold = True
        run = para.add_run(text)
        run.font.size = Pt(10.5)
        return para

    def code(self, text, caption=None):
        if caption:
            cap = self.doc.add_paragraph()
            cap.paragraph_format.space_after = Pt(2)
            cap.paragraph_format.space_before = Pt(8)
            run = cap.add_run(caption)
            run.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = INK2
            keep_with_next(cap)

        for i, line in enumerate(text.rstrip('\n').split('\n')):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(6) if i == 0 else Pt(0)
            pf.left_indent = Cm(0.3)
            pf.line_spacing = 1.0
            run = para.add_run(line if line else ' ')
            run.font.name = 'Consolas'
            run.font.size = Pt(8.8)
            run.font.color.rgb = INK
            shade(para._p.get_or_add_pPr(), SHADE_CODE)
        # trailing spacer
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.space_before = Pt(0)
        for r in spacer.runs:
            r.font.size = Pt(2)

    def callout(self, kind, title, body):
        fills = {'info': SHADE_INFO, 'warn': SHADE_WARN, 'crit': SHADE_CRIT, 'good': SHADE_INFO}
        colors = {'info': ACCENT, 'warn': RGBColor(0x8A, 0x5A, 0x00),
                  'crit': CRITICAL, 'good': GOOD}
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(8)
        para.paragraph_format.space_after = Pt(8)
        para.paragraph_format.left_indent = Cm(0.2)
        run = para.add_run(f'{title}  ')
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = colors[kind]
        body_run = para.add_run(body)
        body_run.font.size = Pt(10)
        body_run.font.color.rgb = INK2
        shade(para._p.get_or_add_pPr(), fills[kind])
        para_borders(para, hex_color='C9C9C2')
        return para

    def table(self, headers, rows, widths=None, font_size=9.5, code_cols=()):
        tbl = self.doc.add_table(rows=1, cols=len(headers))
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tbl.autofit = True

        hdr = tbl.rows[0].cells
        for i, heading in enumerate(headers):
            hdr[i].text = ''
            para = hdr[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.space_before = Pt(2)
            run = para.add_run(str(heading))
            run.bold = True
            run.font.size = Pt(font_size)
            shade(hdr[i]._tc.get_or_add_tcPr(), SHADE_HEAD)

        for row in rows:
            cells = tbl.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ''
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(2)
                run = para.add_run(str(value))
                run.font.size = Pt(font_size)
                if i in code_cols:
                    run.font.name = 'Consolas'
                    run.font.size = Pt(font_size - 0.7)

        if widths:
            for row in tbl.rows:
                for i, width in enumerate(widths):
                    row.cells[i].width = Cm(width)

        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        for r in spacer.runs:
            r.font.size = Pt(2)
        return tbl

    def checklist(self, items, title=None):
        if title:
            self.h4(title)
        self.table(['☐', 'Check', 'How to confirm'],
                   [['', a, b] for a, b in items],
                   widths=[1.0, 6.4, 9.0])

    def save(self, path):
        self.doc.save(str(path))
        return path
