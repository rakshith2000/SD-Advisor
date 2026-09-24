#!/usr/bin/env python3
"""Generates the Aged Ticket Advisor user manual as a Word document.

Audience is Service Desk agents and team leads, not engineers. Kept in the
repo alongside the deployment guide so both regenerate from the same builder
and stay consistent with the application as it changes.

    python docs/generate_user_manual.py
"""

import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from _docx_builder import (ACCENT, CRITICAL, GOOD, INK, INK2, MUTED, Guide)

OUT = Path(__file__).resolve().parent / 'Aged_Ticket_Advisor_User_Manual.docx'
TODAY = datetime.date.today().strftime('%d %B %Y')

g = Guide(footer_label='Aged Ticket Advisor - User Manual',
          header_label='Service Desk - Agents and Team Leads')
d = g.doc


# ---------------------------------------------------------------- cover ----
title = d.add_paragraph()
run = title.add_run('Aged Ticket Advisor')
run.font.size = Pt(32)
run.font.bold = True
run.font.color.rgb = INK

sub = d.add_paragraph()
run = sub.add_run('User Manual')
run.font.size = Pt(18)
run.font.color.rgb = ACCENT

sub2 = d.add_paragraph()
run = sub2.add_run('For Service Desk Agents and Team Leads')
run.font.size = Pt(12.5)
run.font.color.rgb = INK2

d.add_paragraph()
g.p('Your aged tickets, sorted by what actually needs you first — with a suggested '
    'next step and the evidence behind it.', size=11.5, color=INK2)

d.add_paragraph()
g.table(['', ''], [
    ['What it does', 'Reviews every open ticket older than 5 days, works out who is really '
                     'blocking it, ranks them, and suggests what to do next'],
    ['How you use it', 'A daily email, and a web board you can open any time'],
    ['Who it is for', 'Service Desk agents and team leads'],
    ['Does it change my tickets?', 'No. It cannot. It only reads.'],
    ['Do I have to use it?', 'Leads: the morning review moves here. Agents: the digest is '
                             'a help, not an instruction'],
    ['Version / date', f'1.0  ·  {TODAY}'],
], widths=[4.4, 12.0])

d.add_paragraph()
g.callout('info', 'New here? Read three things.',
          'Section 1 (what this is and is not), Section 2.2 (the "blocked by" idea — the one '
          'concept that makes everything else make sense), and then either Section 4 if you '
          'are an agent or Section 5 if you are a lead. Twenty minutes total. Everything else '
          'is reference you can come back to.')


# ------------------------------------------------------------- contents ----
d.add_page_break()
d.add_heading('Contents', 1)
g.p('Right-click the field below in Word and choose "Update Field" to fill in page numbers.',
    italic=True, size=9.5, color=MUTED)

toc_para = d.add_paragraph()
fld_begin = OxmlElement('w:fldChar'); fld_begin.set(qn('w:fldCharType'), 'begin')
instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
instr.text = 'TOC \\o "1-2" \\h \\z \\u'
fld_sep = OxmlElement('w:fldChar'); fld_sep.set(qn('w:fldCharType'), 'separate')
placeholder = OxmlElement('w:t'); placeholder.text = 'Table of contents — press F9 in Word.'
fld_end = OxmlElement('w:fldChar'); fld_end.set(qn('w:fldCharType'), 'end')
r = toc_para.add_run()
for element in (fld_begin, instr, fld_sep, placeholder, fld_end):
    r._r.append(element)

d.add_paragraph()
g.table(['Part', 'Section', 'Who it is for'], [
    ['1', '1. What this is, and what it is not', 'Everyone'],
    ['1', '2. The ideas behind the numbers', 'Everyone'],
    ['1', '3. What the tool suggests, and how much to trust it', 'Everyone'],
    ['2', '4. For agents', 'Agents'],
    ['3', '5. For leads', 'Leads'],
    ['4', '6. Worked examples', 'Everyone'],
    ['4', '7. Frequently asked questions', 'Everyone'],
    ['4', '8. Privacy and your data', 'Everyone'],
    ['4', '9. Glossary', 'Everyone'],
    ['4', '10. Getting help', 'Everyone'],
], widths=[1.6, 10.4, 4.4])


# =========================================================== SECTION 1 =====
g.h1('1. What this is, and what it is not')

g.h2('1.1 The problem it solves')
g.p('Every day, leads work through the aged pending list ticket by ticket: opening each one, '
    'reading the history, working out whether anything has actually happened, and deciding '
    'whether to nudge the agent. It is important work and it takes a long time, and most of '
    'that time is spent on tickets that turn out to be fine.')
g.p('The Aged Ticket Advisor does the reading first. By the time you open the list, every '
    'ticket already has:')
g.bullet('a plain answer to "is anyone actually blocked on us, or are we waiting on someone '
         'else?"')
g.bullet('a score from 0 to 100 saying how much it needs attention, so the list is in the '
         'right order')
g.bullet('a suggested next step, with the reasoning and the evidence behind it')
g.bullet('a ready-to-paste work note and caller message, if one is appropriate')

g.h2('1.2 What it is not')
g.callout('crit', 'It never changes a ticket.',
          'The tool has no ability to write to ServiceNow at all — not a work note, not a '
          'state change, not a reassignment. It reads, it thinks, it tells you. Every action '
          'is still taken by a person, in ServiceNow, exactly as it is today.')

g.table(['People sometimes assume...', 'The reality'], [
    ['"It will close my tickets."',
     'It cannot. It can suggest a ticket looks closable; a person still has to do it.'],
    ['"It is always right."',
     'It is not, and it tells you how confident it is. Below 50% it says so explicitly and '
     'asks you to look properly.'],
    ['"It replaces my judgement."',
     'It does the reading so you can spend your time deciding. The decision is still yours.'],
    ['"It is a performance monitoring tool."',
     'It shows leads the same aged-ticket information they can already pull from ServiceNow, '
     'organised better. See Section 8 for exactly what is visible to whom — we would rather '
     'you knew than guessed.'],
    ['"It reads customer data and sends it to an AI."',
     'It reads ticket data. Phone numbers, email addresses, card and ID numbers, keys and '
     'passwords are masked before anything is sent. Section 8.2 has the detail.'],
], widths=[5.6, 10.8])

g.h2('1.3 What changes for you')
g.table(['', 'Before', 'Now'], [
    ['Lead — morning review',
     'Open the aged list, read each ticket, decide, nudge the agent',
     'Open the digest or the board, work down a ranked list, most tickets need no action'],
    ['Lead — finding the real problems',
     'Spot them by reading everything',
     'They are at the top, flagged, with the reason stated'],
    ['Agent — first you hear of it',
     'A message from your lead',
     'Your own morning email, before your lead asks'],
    ['Agent — writing the update',
     'From scratch',
     'A draft you edit, if you want it'],
    ['Both — disagreeing',
     'A conversation that goes nowhere in particular',
     'One click that is recorded and measured, and changes the tool over time'],
], widths=[3.4, 6.2, 6.8])

g.h2('1.4 Where to find it')
g.table(['Thing', 'Where'], [
    ['The board (all of it)', 'https://<your-server>:8444 — sign in with the account your '
                              'administrator created'],
    ['Your daily email', 'Arrives on weekday mornings from the Service Desk Advisor address'],
    ['A specific ticket', 'Click the ticket number in any email, or go to /board and find it'],
    ['Your own queue', 'On the board, choose your name in the Agent filter'],
], widths=[4.0, 12.4])

g.h2('1.5 A note on the trial period')
g.p('When the tool is first switched on it runs in what the team calls shadow mode. '
    'Everything works — the board, the scoring, the suggestions — but no emails go out. '
    'The point is to let leads check the suggestions against their own judgement for a '
    'couple of weeks before anyone relies on them.')
g.p('If you are reading this during that period, the single most useful thing you can do is '
    'record a verdict on each ticket you look at. Section 5.4 explains how and why.')


# =========================================================== SECTION 2 =====
g.h1('2. The ideas behind the numbers')

g.p('There are four ideas. Once they make sense, everything on the screen makes sense.')

g.h2('2.1 Age and idle time are different things')
g.p('Age is how long the ticket has existed. Idle time is how long since anybody on our side '
    'actually did something to it.')
g.callout('info', 'Idle time is the one that matters.',
          'A ticket can be three weeks old and perfectly healthy — a hardware order with a '
          'known delivery date, say. A ticket that is eight days old with nobody having '
          'touched it for six of them is a different situation entirely.')

g.p('What counts as "somebody actually did something":')
g.table(['Counts as activity', 'Does NOT count'], [
    ['A work note or comment from an agent', 'An automatic system update'],
    ['Changing the state, priority or hold reason', 'An SLA timer recalculating'],
    ['Assigning or reassigning the ticket', 'An integration account touching the record'],
    ['Attaching a knowledge article', 'The caller adding a comment (that is their activity, '
                                      'not ours — and it usually means we now owe a reply)'],
], widths=[8.2, 8.2])

g.h2('2.2 Blocked by — the important one')
g.p('Every aged ticket gets a "Blocked by" value. This is the tool answering the question '
    'leads actually care about: is this one ours to move, or are we legitimately waiting?')

g.table(['Blocked by', 'What it means', 'Is it ours?'], [
    ['Us', 'Nothing is stopping us. We owe the next action', 'Yes — act'],
    ['Caller', 'We are waiting on information or confirmation from the caller', 'No — but see below'],
    ['Vendor', 'A third party is holding it up', 'No — but chase if it has been a while'],
    ['Change', 'Waiting on a change request', 'No'],
    ['Problem', 'Waiting on a problem record', 'No'],
    ['Approval', 'Waiting on someone to approve something', 'No'],
], widths=[2.6, 9.4, 4.4])

g.callout('good', 'Why this saves the most time.',
          'A list of forty aged tickets is daunting. "Nine of these are ours and thirty-one '
          'are legitimately waiting on somebody else" is a morning\'s work you can actually '
          'do. That single reframing is the biggest thing the tool gives you.')

g.h3('Two cases where the tool overrules the ticket')
g.p('The Blocked by value is not simply copied from the hold reason. There are two situations '
    'where the ticket says one thing and reality says another:')
g.table(['Situation', 'What the ticket says', 'What the tool says', 'Why'], [
    ['The caller has replied',
     'On Hold — Awaiting Caller',
     'Blocked by: Us',
     'They answered. The hold reason is just stale. We owe them a response'],
    ['The blocker has closed',
     'On Hold — Awaiting Change',
     'Blocked by: Us',
     'That change went in last Tuesday. Nothing is blocking this ticket any more'],
], widths=[3.4, 4.0, 3.0, 6.0])
g.p('These two cases are, in practice, where most genuinely stuck tickets hide. Nobody is at '
    'fault — there is simply no mechanism in ServiceNow that tells you a dependency cleared.')

g.h2('2.3 The attention score')
g.p('A number from 0 to 100 that puts the list in the right order. Higher means "look at this '
    'one sooner". It is arithmetic, not opinion — and the ticket page shows you exactly how '
    'each ticket earned its number.')

g.table(['Score', 'Label', 'What it usually means'], [
    ['70–100', 'Critical', 'Breached or about to breach, and stalled. Deal with it today'],
    ['45–69', 'High', 'Something is genuinely wrong here. Deal with it this shift'],
    ['25–44', 'Medium', 'Worth a look, not urgent'],
    ['0–24', 'Low', 'Aged, but nothing is actually wrong. Usually needs nothing'],
], widths=[2.2, 2.4, 11.8])

g.p('What feeds the score, and roughly how much each contributes:')
g.table(['Ingredient', 'Weight', 'Plain meaning'], [
    ['SLA jeopardy', '25', 'How close it is to breaching, or whether it already has'],
    ['Stagnation', '25', 'How long since anybody touched it'],
    ['Blocked but stale', '15', 'Caller waiting on a reply, blocker already cleared, or never '
                                'actioned at all'],
    ['Age', '10', 'How far past the 5-day threshold. Deliberately a small ingredient — age on '
                  'its own means very little'],
    ['Priority', '10', 'P1 outranks P4'],
    ['Past expected time', '10', 'Slow compared with other tickets of the same type, not slow '
                                 'in the abstract'],
    ['Reassignment churn', '5', 'Bounced between queues, or reopened'],
], widths=[3.4, 1.6, 11.4])

g.callout('info', 'Slow "for its own type".',
          'The tool learns how long tickets in each category normally take by looking at ones '
          'you have already resolved. A printer request that usually takes two days is flagged '
          'at four; a laptop build that normally takes ten days is not. This stops slow-by-'
          'nature work sitting permanently red.')

g.h2('2.4 Risk flags')
g.p('Short labels on each ticket saying what specifically is wrong. These are the things to '
    'scan for.')
g.table(['Flag', 'What it means', 'Typical response'], [
    ['SLA breached', 'The resolution SLA has already been missed',
     'Prioritise, and make sure the reason is documented'],
    ['SLA at risk', 'Past 75% of the allowed time',
     'Act now while it is still avoidable'],
    ['No action 4+ days', 'Nobody on our side has touched it for four days or more',
     'Either progress it or record why it is waiting'],
    ['No recent action', 'Untouched for two days or more',
     'A nudge, usually'],
    ['Caller awaiting reply', 'The caller responded and has had no answer since',
     'Reply. This is the most avoidable kind of delay there is'],
    ['Blocker already closed', 'The change or problem it was waiting for has closed',
     'Take it off hold and carry on'],
    ['Closure candidate', 'Follow-ups documented, caller silent since',
     'Consider closing, per your local policy'],
    ['Past expected time', 'Slower than tickets of the same type usually take',
     'Worth understanding why'],
    ['KB not attached', 'A knowledge article covers this and was not linked',
     'Attach it — helps the next person and the caller'],
    ['Never actioned', 'No agent has touched it since it was raised',
     'Pick it up, or get it assigned properly'],
    ['Suspicious content', 'Something in the ticket text looked like an attempt to manipulate '
                           'the tool',
     'Rare. Mention it to your lead'],
], widths=[3.2, 7.0, 6.2])


# =========================================================== SECTION 3 =====
g.h1('3. What the tool suggests, and how much to trust it')

g.h2('3.1 The eight suggestions')
g.p('Every analysed ticket gets exactly one suggested next step.')
g.table(['Suggestion', 'Means', 'Usually right when'], [
    ['Resolve now', 'There is enough here to resolve it',
     'The fix is documented, or the blocker has cleared'],
    ['Follow up with caller', 'We owe the caller contact',
     'They replied and got nothing back, or we are waiting on them and have gone quiet'],
    ['Chase vendor', 'A third party is holding it and needs pushing',
     'No vendor update for a while'],
    ['Reassign', 'It is in the wrong queue',
     'The evidence points clearly at a specific named team'],
    ['Escalate', 'Needs a lead or someone senior',
     'Stuck a long time, or the tool is not confident enough to be more specific'],
    ['Waiting on dependency', 'Genuinely blocked, correctly',
     'The change or problem it is waiting on is still open'],
    ['Close (no response)', 'Follow-ups exhausted, caller silent',
     'Three or more documented attempts and nothing back'],
    ['On track', 'Nothing needs doing',
     'It is aged but progressing, or legitimately waiting'],
], widths=[3.2, 5.4, 7.8])

g.callout('good', '"On track" is a real answer, not a cop-out.',
          'A tool that finds something wrong with every ticket gets ignored within a fortnight. '
          'This one is explicitly allowed — and expected — to say a ticket is fine. If most of '
          'your aged list comes back On track, that is good news about your queue, not a '
          'broken tool.')

g.h2('3.2 Confidence')
g.p('Every suggestion carries a confidence percentage.')
g.table(['Confidence', 'How to read it'], [
    ['Above 80%', 'The evidence is clear. Usually safe to act on after a sanity check'],
    ['50–80%', 'Reasonable, but read the reasoning before you act'],
    ['Below 50%', 'Flagged as "Needs your review". The tool is telling you it cannot judge '
                  'this one properly — usually because the ticket history is too thin. '
                  'Treat it as a prompt to look, not as advice'],
], widths=[3.0, 13.4])
g.p('Low confidence is a feature. The alternative — a confident-sounding guess — is worse '
    'than an honest "I am not sure about this one".')

g.h2('3.3 Evidence')
g.p('Suggestions are not produced from thin air. Before deciding, the tool looks up:')
g.bullet('tickets that look like this one and have already been resolved, including what '
         'actually fixed them and how long they took;', bold_prefix='Similar incidents — ')
g.bullet('published articles relevant to the symptom, and whether one is already attached.',
         bold_prefix='Knowledge articles — ')
g.p('The ticket page lists exactly which ones it used, with their numbers, so you can open '
    'them and check. If the reasoning cites INC2811111, you can go and read INC2811111.')

g.h2('3.4 The drafted text')
g.p('Where it makes sense, the tool writes two things for you:')
g.table(['Draft', 'Written as', 'Use it for'], [
    ['Work note', 'Agent to record — factual, past tense, no greeting',
     'Documenting what is happening, internally'],
    ['Message to caller', 'Agent to caller — courteous, plain language, no jargon or hostnames',
     'The chase or update you were going to have to write anyway'],
], widths=[3.0, 6.6, 6.8])
g.callout('warn', 'Always read it before you send it.',
          'It is a first draft written by a machine that has read the ticket. It is usually a '
          'good starting point and it is sometimes subtly wrong. Your name goes on it, not the '
          'tool\'s. Edit freely.')

g.h2('3.5 When it is wrong')
g.p('It will be, sometimes. What matters is what happens next:')
g.numbered('Do the right thing for the ticket. Your judgement wins, always.')
g.numbered('Tell the tool. On the ticket page, record a verdict of "Wrong" or "Not '
           'applicable" and, if you have ten seconds, say why in the comment box.')
g.numbered('That verdict is stored and reported. If the same kind of mistake shows up '
           'repeatedly, the pattern is visible on the Accuracy page and the configuration '
           'gets fixed.')
g.p('Ignoring a bad suggestion is understandable. Recording it is what actually makes it '
    'stop happening.')


# =========================================================== SECTION 4 =====
g.h1('4. For agents')

g.h2('4.1 Your five-minute morning')
g.numbered('Open the email titled "Your aged tickets".')
g.numbered('Look at the second number in the summary line — how many are waiting on you. '
           'That is your actual list.')
g.numbered('Work down them. Each one tells you what it thinks you should do and why.')
g.numbered('Where a draft is offered and it fits, copy it, edit it, use it.')
g.numbered('Ignore anything marked On track. It is there for completeness.')

g.h2('4.2 Your daily email, explained')
g.p('The email has three parts.')

g.h3('The summary line')
g.code("""14 aged  ·  5 waiting on you  ·  2 SLA breached  ·  1 caller awaiting your reply""")
g.p('Only the second number is your to-do list. The other nine tickets are waiting on a '
    'caller, a vendor or a change, and are not yours to move today.')

g.h3('Each ticket')
g.p('A coloured square and a label (Critical, High, Medium, Low), the ticket number and '
    'summary, how old it is and how long it has been idle, who is blocking it, and any risk '
    'flags. Then a box with the suggested next step and the reasoning.')

g.h3('Ready to paste')
g.p('Where a draft work note or caller message exists, it appears at the bottom of that '
    'ticket. Copy, edit, send.')

g.h2('4.3 Using the board')
g.p('The email is a snapshot of the morning. The board is live and you can filter it.')
g.table(['To see', 'Do this'], [
    ['Only your tickets', 'Agent filter → your name'],
    ['Only what is actually yours to move', 'Blocked by → Us'],
    ['Only the urgent ones', 'Min score → 45'],
    ['Only SLA problems', 'Risk flag → SLA breached, or SLA at risk'],
    ['Everything you could close', 'Suggested step → Close (no response)'],
], widths=[5.6, 10.8])

g.h2('4.4 The ticket page')
g.p('Click any ticket number. The page has, left to right:')
g.table(['Area', 'Contains'], [
    ['Suggested next step', 'The recommendation, the reasoning, a confidence bar, and the '
                            'specific tickets and articles it used as evidence'],
    ['Ready to paste', 'The drafted work note and caller message, each with a Copy button'],
    ['Activity', 'The ticket history with the system noise stripped out — just the things '
                 'people actually did, newest first, colour-coded by whether it was an agent '
                 'or the caller'],
    ['Signals (sidebar)', 'Age, idle time, last agent action, last caller activity, SLA state, '
                          'follow-ups sent, dependency, KB status'],
    ['Why it ranks here', 'The score broken into its ingredients, so you can see exactly what '
                          'pushed it up the list'],
    ['Open in ServiceNow', 'A direct link, because all the actual work still happens there'],
], widths=[3.6, 12.8])

g.h2('4.5 A playbook for each suggestion')
g.table(['If it says', 'Do this'], [
    ['Follow up with caller',
     'Check the Activity panel for whether they already replied. If they did, answer that '
     'first — do not send a generic chase to someone who is waiting on you'],
    ['Chase vendor',
     'Update the work note with the vendor reference and what you asked for. That is what '
     'stops it being flagged again tomorrow'],
    ['Reassign',
     'Check the suggested group makes sense. The tool can only suggest groups that really '
     'exist, but it can still pick the wrong one. Put the reason in a work note'],
    ['Resolve now',
     'Read the similar incident it cites. If that fix applies, apply it, and attach the KB '
     'article while you are there'],
    ['Waiting on dependency',
     'Nothing to do. Check the dependency reference in the sidebar is still open — if it has '
     'closed, the flag will already have changed to "Blocker already closed"'],
    ['Close (no response)',
     'Follow your local closure policy. The tool has checked that the follow-ups are '
     'documented; it is not authorising the closure'],
    ['Escalate', 'Flag it to your lead with what you have tried'],
    ['On track', 'Nothing'],
], widths=[3.4, 13.0])

g.h2('4.6 Habits that keep your tickets off the list')
g.p('Not a scoring system to game — these are the same habits that make tickets easy for the '
    'next person to pick up.')
g.table(['Habit', 'Why it helps'], [
    ['Put a work note on when you do something',
     'The idle clock resets on real activity. Work you did not record did not happen, as far '
     'as any tool or any lead can tell'],
    ['Use the right hold reason',
     'This is what drives Blocked by. "Awaiting Vendor" on a ticket that is really waiting on '
     'the caller sends it to the wrong place'],
    ['Make chase messages look like chases',
     'Phrases like "following up", "gentle reminder", "any update" are how documented '
     'follow-ups are counted. Three counted follow-ups is what makes a ticket eligible for '
     'closure instead of sitting open forever'],
    ['Attach the KB article you used',
     'Clears the "KB not attached" flag, and genuinely helps the caller and the next agent'],
    ['Take it off hold when the blocker clears',
     'The tool will catch it if you do not, but a ticket sitting on hold for a change that '
     'closed two weeks ago is pure lost time'],
], widths=[5.0, 11.4])

g.h2('4.7 If a ticket is flagged unfairly')
g.p('The most common cause is that real work was done but not recorded in a way anything can '
    'see — a phone call with no work note, or an update made by an automated account that the '
    'tool has been told to ignore.')
g.p('Tell your lead. If it is a one-off, they record "Not applicable" and it goes away. If it '
    'keeps happening to the same kind of ticket, it is a configuration problem and worth '
    'fixing properly — Section 5.8.')


# =========================================================== SECTION 5 =====
g.h1('5. For leads')

g.h2('5.1 Your ten-minute morning review')
g.numbered('Open the digest, or go straight to the board.')
g.numbered('Read the "Since yesterday" line. New and worsening tickets are the only ones that '
           'have actually changed since you last looked.')
g.numbered('Work down the ranked list from the top. Stop when the scores drop below about 45 '
           '— below that, nothing is usually wrong.')
g.numbered('On each one, record a verdict. One click.')
g.numbered('Check the Closure candidates block and deal with it as a batch.')
g.numbered('Once a week, open By agent and Accuracy.')

g.callout('good', 'The point is that you stop reading tickets that are fine.',
          'On a typical day the top five to ten entries are the whole job. Everything below '
          'them has already been read and found to be progressing or legitimately waiting.')

g.h2('5.2 Your digest, explained')
g.table(['Block', 'What it is for'], [
    ['Eight summary tiles', 'The shape of the backlog at a glance. "Waiting on us" is the one '
                            'that matters — it is your real workload'],
    ['Since yesterday', 'Newly aged and newly worse. If you are short of time, read only this'],
    ['Ranked list', 'Up to 25 tickets, worst first, each with the suggestion and why it ranks '
                    'where it does'],
    ['Closure candidates', 'Follow-ups exhausted and caller silent. Handle as a batch rather '
                           'than one at a time'],
    ['By agent', 'Who is holding what, worst first, with their audit compliance average if '
                 'that link is enabled'],
], widths=[3.6, 12.8])

g.h2('5.3 The board')
g.p('Everything in the digest, live and filterable. The filters that get the most use:')
g.table(['Filter', 'Use it to answer'], [
    ['Blocked by = Us', '"What is genuinely ours this morning?"'],
    ['Risk flag = Caller awaiting reply', '"Who are we leaving hanging?"'],
    ['Risk flag = Blocker already closed', '"What is stuck for no reason at all?"'],
    ['Suggested step = Close (no response)', '"What can we clear out?"'],
    ['Agent = <name>', 'Prepare for a one-to-one'],
    ['Min score = 70', '"Just show me the fires"'],
], widths=[5.6, 10.8])

g.h2('5.4 Recording verdicts — the most important thing you do')
g.callout('crit', 'Without verdicts there are no accuracy numbers.',
          'The claim "this tool is 84% accurate" is only meaningful because leads pressed '
          'buttons. If nobody records verdicts, nobody can tell whether the suggestions are '
          'any good, and the honest answer to "should we trust it?" stays "we do not know".')

g.p('On any ticket page, the Your verdict panel offers five options:')
g.table(['Verdict', 'Use when', 'Side effect'], [
    ['Agree', 'The suggested action is what you would have done', 'Counts towards agreement'],
    ['Mostly right — I adjusted it', 'Right direction, wrong detail', 'Counts towards agreement'],
    ['Wrong', 'Not the right action for this ticket', 'Counts against. This is the useful one'],
    ['Not applicable', 'This should not have been flagged at all', 'Counts against precision, '
                                                                   'and snoozes it for 48 hours'],
    ['Already handled', 'You dealt with it before you saw the suggestion',
     'Excluded from precision, and snoozes it for 48 hours'],
], widths=[3.4, 6.0, 7.0])

g.p('The optional "what you actually did" dropdown is worth using when you disagree — it is '
    'the difference between knowing the tool was wrong and knowing what it should have said.')

g.h2('5.5 Snoozing')
g.p('If you have dealt with a ticket and do not want to see it again for a while, snooze it '
    'for 24 hours, 2 days, 3 days or a week. Snoozed tickets disappear from the board and the '
    'digests until the time expires.')
g.callout('info', 'Use it freely.',
          'A tool that keeps raising a ticket you have already handled trains you to ignore it. '
          'Snoozing is not hiding a problem; it is telling the tool you have this one. You can '
          'always see snoozed tickets with the "Include snoozed" checkbox.')

g.h2('5.6 The By agent page')
g.p('Per-agent rollup, sorted worst-first by their highest-scoring ticket. Columns: aged '
    'count, worst score, average idle days, SLA breaches, tickets stale four days or more, '
    'callers waiting, and KB gaps.')
g.p('If the link to the ticket audit tool is enabled, there is also a 30-day compliance '
    'average from that system. The combination is genuinely new information: "nine aged '
    'tickets" is a workload observation, but "nine aged tickets, all stalled on hold without '
    'follow-up, alongside a 52% closure-note compliance average" is a specific, evidenced '
    'coaching conversation.')

g.callout('warn', 'Use it as a conversation starter, not a league table.',
          'High numbers frequently mean a hard queue, a period of leave, or an unlucky '
          'allocation rather than poor performance. Open the actual tickets before drawing a '
          'conclusion — the page links straight to them. Treating this as a ranking is the '
          'fastest way to lose the team\'s trust in the whole tool.')

g.h2('5.7 The Accuracy page')
g.p('Four numbers, all derived from the verdicts you record.')
g.table(['Metric', 'Question it answers', 'Target'], [
    ['Action agreement', 'When we suggested something, was it the right thing?', '80% or better'],
    ['Actionable precision', 'When we flagged a ticket, did it deserve flagging?', '85% or better'],
    ['Review coverage', 'How many suggestions were actually judged? Low coverage means the '
                        'first two numbers come from a small, self-selected sample',
     'As high as you can manage'],
    ['Silent breaches', 'Tickets that breached with no prior warning. The opposite failure — '
                        'not a wrong alert, a missing one', '0'],
], widths=[3.4, 9.0, 4.0])

g.p('The Agreement by suggested action table is where the useful detail is. It is common for '
    'a tool like this to be reliable on some kinds of decision and weak on others. If '
    '"Follow up with caller" agrees 90% of the time and "Reassign" agrees 55%, then trust the '
    'first and check the second — and raise the second with whoever maintains the tool.')

g.h2('5.8 Tuning')
g.p('Administrators can adjust how the score is calculated at Settings → Weights. The seven '
    'ingredients must total 100.')
g.table(['If leads say', 'Consider'], [
    ['"It flags too much"', 'Raise the aged threshold from 5 days, or narrow the groups in scope'],
    ['"The order feels wrong"', 'Raise Stagnation if untouched tickets matter most, or SLA '
                                'jeopardy if breach risk dominates'],
    ['"It misses tickets we care about"', 'Lower the threshold, or check the group filter is '
                                          'not too narrow'],
    ['"Idle times are nonsense"', 'An automated account is missing from the configuration. '
                                  'This is the most common problem and needs an administrator'],
    ['"Blocked by is often wrong"', 'Your hold reasons may not match the built-in mapping. '
                                    'Needs an administrator'],
], widths=[5.0, 11.4])

g.h2('5.9 Alerts between digests')
g.p('Some things should not wait until tomorrow morning. Leads get an immediate batched email '
    'when a ticket newly breaches SLA, newly passes 75% of its SLA, becomes critically stale, '
    'has a caller waiting on a reply, or has its blocker clear.')
g.p('Each ticket alerts at most once per day per type, so a ticket that stays breached does '
    'not email you every ten minutes.')


# =========================================================== SECTION 6 =====
g.h1('6. Worked examples')

def scenario(n, title, board, tool_says, whats_going_on, do_this):
    g.h2(f'6.{n} {title}')
    g.p('What you see on the board', bold=True, after=2)
    g.code(board)
    g.rich([('The tool says:  ', {'b': True}), (tool_says, {})], after=4)
    g.rich([('What is going on:  ', {'b': True}), (whats_going_on, {})], after=4)
    g.rich([('What to do:  ', {'b': True}), (do_this, {})])

scenario(1, 'The blocker cleared and nobody noticed',
         """INC2831044   score 74  Critical
Printer driver deployment failing on the Chicago floor
Asha Rao · 19 days old · idle 11.4 days · Blocked by: Us
Flags: Blocker already closed · No action 4+ days · Past expected time""",
         '"Resolve now" — confidence 86%. The change this ticket was held for, CHG0044321, '
         'closed successfully eleven days ago.',
         'The ticket is still On Hold — Awaiting Change. Nobody is at fault: ServiceNow does '
         'not tell you when a dependency closes, so the ticket simply sat there. The tool '
         'noticed the change state and moved the ball back to us.',
         'Take it off hold, confirm the driver deployment worked, resolve it. This category of '
         'ticket is often the oldest thing on the board and the quickest to clear.')

scenario(2, 'The caller replied four days ago',
         """INC2838464   score 81  Critical
Cannot connect to VPN from home
Priya Sharma · 9 days old · idle 4.2 days · Blocked by: Us
Flags: SLA breached · Caller awaiting reply · No action 4+ days""",
         '"Follow up with caller" — confidence 88%. The caller supplied the requested error '
         'screenshot on 19 September and has had no response since.',
         'The ticket is On Hold — Awaiting Caller, so it drops out of the normal "things I '
         'need to do" view. But the caller answered. The hold reason is just stale, and the '
         'SLA clock kept running.',
         'Reply to the caller today. This is the single most avoidable category of delay in '
         'any service desk, and the easiest for the tool to catch.')

scenario(3, 'Ready to close, nobody pulled the trigger',
         """INC2836210   score 38  Medium
Request for additional monitor
Dev Menon · 22 days old · idle 6.1 days · Blocked by: Caller
Flags: Closure candidate · Past expected time""",
         '"Close (no response)" — confidence 79%. Three documented follow-ups on 2, 8 and 14 '
         'September. No caller response for eighteen days.',
         'The agent did everything right and documented all of it. The ticket is just sitting '
         'there because nobody wants to be the one to close it.',
         'Use the Closure candidates block in the digest and handle these as a batch, per your '
         'local closure policy. A drafted closure note is provided.')

scenario(4, 'The tool gets it wrong',
         """INC2837788   score 52  High
Application error on month-end reporting run
Sam Okafor · 12 days old · idle 3.0 days · Blocked by: Us
Flags: No recent action · KB not attached""",
         '"Reassign to Network Ops" — confidence 61%.',
         'Reading the work notes, this is a finance application defect that the vendor is '
         'already working under a support case. The tool latched onto a mention of a network '
         'timeout in an early work note and drew the wrong conclusion.',
         'Do the right thing — leave it with the vendor. Then record the verdict as "Wrong", '
         'set "what you actually did" to Chase vendor, and add one line to the comment box. '
         'That is how the pattern gets spotted and fixed.')

scenario(5, 'Old, and completely fine',
         """INC2835901   score 16  Low
New starter laptop build - Shanghai
Wei Chen · 14 days old · idle 0.8 days · Blocked by: Vendor
Flags: (none)""",
         '"On track" — confidence 84%. The hardware order is with the supplier, the reference '
         'is recorded, and the ETA is documented in the latest work note.',
         'A two-week-old ticket that is entirely healthy. The agent is on top of it, and '
         'laptop builds in this category normally take around this long.',
         'Nothing. This is the tool doing its most valuable work — telling you that a ticket '
         'which would have caught your eye on an age-sorted list does not need you.')

scenario(6, 'Nobody ever picked it up',
         """INC2838120   score 69  High
Shared mailbox permissions request
Unassigned · 8 days old · idle 8.0 days · Blocked by: Us
Flags: Never actioned · SLA at risk · No action 4+ days""",
         '"Escalate" — confidence 58%. No agent activity of any kind since the ticket was '
         'raised.',
         'It was never assigned. It has been in the queue for eight days with nobody looking '
         'at it. Idle time runs from creation, because there is no agent action to measure '
         'from.',
         'Assign it now. Then ask why it was missed — this is usually a queue routing or '
         'triage issue rather than anything to do with an individual agent.')


# =========================================================== SECTION 7 =====
g.h1('7. Frequently asked questions')

g.h2('7.1 About the suggestions')
faqs_general = [
    ('Can it change or close my tickets?',
     'No. It has no ability to write to ServiceNow at all. Every action is still taken by a '
     'person.'),
    ('What if I disagree with a suggestion?',
     'Do the right thing for the ticket. Then record a verdict of Wrong or Not applicable so '
     'the disagreement is counted. That is how the tool improves.'),
    ('Why is a ticket I am actively working being flagged?',
     'Almost always because the work is not visible in the ticket — a phone call with no work '
     'note, for example. Add the note. If the work IS recorded and it is still flagged, tell '
     'your lead: an automated account may be missing from the configuration.'),
    ('The suggested work note is wrong. Is that a bug?',
     'Not necessarily. It is a first draft written from the ticket history. Edit it before '
     'you use it — that is what it is for. If it is badly wrong, record the verdict.'),
    ('Does it work on tickets that are not in English?',
     'It reads and reasons over them, and your queue includes plenty. Quality there has been '
     'tested less thoroughly than for English tickets, so please flag anything that looks '
     'poor.'),
    ('Why does it say "On track" for a ticket that is three weeks old?',
     'Because age alone does not mean something is wrong. If it is progressing or legitimately '
     'waiting, that is the correct answer.'),
    ('It suggested reassigning to a group that does not exist.',
     'It should not be able to — suggested groups are checked against the real list, and an '
     'unrecognised one is downgraded to Escalate. If you see this, report it.'),
]
g.table(['Question', 'Answer'], faqs_general, widths=[5.0, 11.4])

g.h2('7.2 About how it works')
faqs_how = [
    ('How often does it update?',
     'Ticket data and the scores refresh every ten minutes. The suggestions refresh hourly, '
     'and always before a digest goes out.'),
    ('Why did the score change overnight when nothing happened?',
     'Idle time and age keep growing, so a ticket nobody touches slowly climbs. That is '
     'intentional.'),
    ('Where does "tickets like this normally take 4 hours" come from?',
     'From your own resolved tickets in the same category over the last six months — not an '
     'industry benchmark.'),
    ('Why do some tickets have no suggestion?',
     'They have been synced and scored but not yet analysed. They pick one up within the hour.'),
    ('What happens if the AI service is down?',
     'The board, the scores and the flags keep working — they are ordinary arithmetic and do '
     'not need it. Suggestions fall back to simple rules and say so.'),
    ('Does it read attachments?',
     'The advisor reads ticket fields, history, SLA records and knowledge articles. It does '
     'not open attachments.'),
]
g.table(['Question', 'Answer'], faqs_how, widths=[5.0, 11.4])

g.h2('7.3 The questions people ask but do not always say out loud')
g.callout('info', 'These are fair questions and deserve straight answers.', '')
faqs_hard = [
    ('Is this monitoring my performance?',
     'It shows leads aged-ticket information they can already get from ServiceNow today, '
     'organised more usefully. What is genuinely new is the per-agent rollup on one page, and '
     'optionally a compliance average from the existing ticket audit tool. So: it does not '
     'create new surveillance, but it does make existing information easier to see. You should '
     'know that rather than find out later.'),
    ('Will my manager see a score for me?',
     'There is no score for a person. The By agent page shows counts — how many aged tickets, '
     'average idle days, SLA breaches. If the audit link is switched on it also shows your '
     '30-day compliance average, which already exists in the audit tool.'),
    ('Could this be used against me?',
     'The manual explicitly tells leads to treat the page as a conversation starter and to '
     'open the actual tickets before drawing conclusions, because high numbers usually mean a '
     'hard queue or time off rather than poor work. If it is being used as a league table, '
     'that is a management issue worth raising.'),
    ('Is it going to replace service desk jobs?',
     'It does not resolve tickets, talk to callers or make decisions. It removes the reading '
     'that leads currently do before they can decide anything.'),
    ('Is my writing being judged by an AI?',
     'The tool reads work notes to work out what happened and whether follow-ups were made. '
     'It produces one coaching line per ticket aimed at the lead — about handling, not about '
     'writing style.'),
    ('What if I just ignore it?',
     'Nothing happens automatically. But the digest is usually faster than the conversation '
     'you would otherwise have with your lead, so it is generally worth the five minutes.'),
]
g.table(['Question', 'Answer'], faqs_hard, widths=[5.0, 11.4])


# =========================================================== SECTION 8 =====
g.h1('8. Privacy and your data')

g.h2('8.1 What the tool reads')
g.table(['Source', 'What it reads'], [
    ['Incidents', 'Standard fields — description, category, priority, state, hold reason, '
                  'assignment, caller, timestamps'],
    ['Ticket history', 'Who changed what and when, including work notes and comments'],
    ['SLA records', 'Breach state, percentage consumed, planned end time'],
    ['Knowledge articles', 'Published articles, to find relevant ones'],
    ['Change and problem records', 'Only the number and state of a linked record, to see '
                                   'whether it has closed'],
    ['Resolved incidents', 'Closure notes from past tickets, to learn what fixes what'],
], widths=[4.4, 12.0])
g.p('All of this is data that leads can already see in ServiceNow. Nothing hidden is being '
    'surfaced.')

g.h2('8.2 What is hidden before anything reaches the AI')
g.p('Ticket text is scrubbed before it is sent to the language model. Replaced with '
    'placeholders:')
g.table(['Masked', 'Examples'], [
    ['Email addresses', 'jamie.fox@example.com'],
    ['Phone numbers', 'Any recognisable phone or mobile number'],
    ['Payment and identity numbers', 'Card numbers, national ID and social security formats, '
                                     'IBANs'],
    ['Credentials', 'Anything written as "password is X", "token: Y" and similar'],
    ['Keys and tokens', 'API keys, JWTs, private key blocks'],
    ['IP addresses', 'Internal addressing'],
], widths=[5.0, 11.4])
g.p('Placeholders look like [[EMAIL_1]] or [[PHONE_2]]. The real values are put back before '
    'anything is shown to you, so a drafted message still reads properly — but the model '
    'itself never saw them.')

g.h2('8.3 What is stored, and for how long')
g.table(['Stored', 'Retention'], [
    ['Ticket snapshot and computed signals', '180 days of history, then automatically deleted'],
    ['Suggestions and their reasoning', '180 days once superseded'],
    ['Verdicts leads record', 'Kept — this is the accuracy record'],
    ['Alert history', '180 days'],
    ['Email delivery log', 'Kept'],
], widths=[6.0, 10.4])
g.p('Everything lives in a database on your own infrastructure. The only thing that leaves '
    'the estate is the masked text sent to your organisation\'s own Azure OpenAI service.')

g.h2('8.4 Who can see what')
g.table(['Role', 'Can see', 'Can do'], [
    ['Viewer', 'The board, ticket pages, agent rollup, accuracy', 'Nothing — read only'],
    ['Lead', 'The same', 'Record verdicts, snooze, request re-analysis'],
    ['Admin', 'The same', 'All lead actions, plus tune the score and send digests manually'],
], widths=[2.4, 7.0, 7.0])
g.p('Leads may be restricted to specific assignment groups, in which case they see only those '
    'groups on the board and in their digest.')


# =========================================================== SECTION 9 =====
g.h1('9. Glossary')
g.table(['Term', 'Meaning'], [
    ['Aged ticket', 'An open ticket older than the configured threshold, normally 5 days'],
    ['Attention score', 'The 0–100 ranking number. Higher means look sooner'],
    ['Blocked by', 'Who owes the next action: Us, Caller, Vendor, Change, Problem or Approval'],
    ['Closure candidate', 'Follow-ups documented and caller silent long enough that closure is '
                          'worth considering'],
    ['Confidence', 'How sure the tool is about its suggestion. Below 50% is flagged as needing '
                   'your review'],
    ['Digest', 'The daily email. Leads get the whole queue; agents get their own tickets'],
    ['Evidence', 'The specific past tickets and knowledge articles the suggestion was based on'],
    ['Idle time', 'Days since anyone on our side last did something real to the ticket'],
    ['Risk flag', 'A short label naming a specific problem, such as "Caller awaiting reply"'],
    ['Shadow mode', 'The trial period. Everything runs but no emails are sent'],
    ['Silent breach', 'A ticket that breached SLA without ever being flagged as at risk first'],
    ['Snooze', 'Hide a ticket you have dealt with, for 24 hours up to a week'],
    ['Verdict', 'Your one-click judgement on whether a suggestion was right'],
], widths=[3.6, 12.8])


# ========================================================== SECTION 10 =====
g.h1('10. Getting help')

g.h2('10.1 Who to ask')
g.table(['Issue', 'Who'], [
    ['A specific suggestion looks wrong', 'Record the verdict. No need to raise anything'],
    ['The same kind of mistake keeps happening', 'Your team lead, who raises it with the tool '
                                                 'owner'],
    ['Idle times or Blocked by are consistently wrong', 'Tool owner — this is a configuration '
                                                        'fix, not a one-off'],
    ['Cannot sign in, or the board will not load', 'Tool owner or platform support'],
    ['Not receiving the digest', 'Tool owner — they can check the delivery log and tell you '
                                 'whether it was sent'],
    ['A ticket is missing from the board entirely', 'Tool owner — usually an assignment group '
                                                    'that is not in scope'],
], widths=[6.4, 10.0])

g.h2('10.2 Useful things to include when reporting a problem')
g.bullet('The ticket number.')
g.bullet('What the tool suggested, and what the right answer was.')
g.bullet('Roughly when you looked at it — suggestions change as tickets change.')
g.bullet('A screenshot of the ticket page, if you can.')

g.h2('10.3 One thing to remember')
g.callout('good', 'Your judgement outranks the tool, every time.',
          'It is there to do the reading so that you have more time to decide, and to make '
          'sure nothing quietly rots at the bottom of a queue. It is not there to tell you '
          'your job. When it is right, it saves you twenty minutes. When it is wrong, press '
          'the button that says so — and it gets better.')


# ---------------------------------------------------------------------------
if __name__ == '__main__':
    OUT.parent.mkdir(parents=True, exist_ok=True)
    g.save(OUT)
    print(f'Written: {OUT}')
    print(f'Size:    {OUT.stat().st_size / 1024:.0f} KB')
    sys.exit(0)
